# -*- coding: utf-8 -*-
"""生成 Office 测试文件，用于验证 office.js 的 Markdown 转换。"""
import os
from pptx import Presentation
from pptx.util import Inches
from docx import Document
from openpyxl import Workbook

out = os.path.join(os.environ.get("TEMP", "."), "office-test")
os.makedirs(out, exist_ok=True)

# ---- PPTX ----
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "光的折射"
body = slide.placeholders[1]
body.text = "光从空气斜射入水中时，折射角小于入射角。"
tf = body.text_frame
p = tf.add_paragraph()
p.text = "入射角增大，折射角也增大。"
p2 = tf.add_paragraph()
p2.text = "易错点：垂直入射时方向不变，折射角为 0 度。"

slide2 = prs.slides.add_slide(prs.slide_layouts[1])
slide2.shapes.title.text = "凸透镜成像"
slide2.placeholders[1].text = "物距大于二倍焦距时，成倒立缩小的实像。"
notes = slide2.notes_slide.notes_text_frame
notes.text = "课堂备注：结合照相机原理讲解。"
prs.save(os.path.join(out, "test.pptx"))

# ---- DOCX ----
doc = Document()
doc.add_heading("一元二次方程复习", level=1)
doc.add_paragraph("一般形式：ax² + bx + c = 0（a ≠ 0）。")
doc.add_heading("求根公式", level=2)
doc.add_paragraph("x = (-b ± √(b² - 4ac)) / (2a)")
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "判别式"
table.cell(0, 1).text = "根的情况"
table.cell(1, 0).text = "Δ > 0"
table.cell(1, 1).text = "两个不等实根"
doc.save(os.path.join(out, "test.docx"))

# ---- XLSX ----
wb = Workbook()
ws = wb.active
ws.title = "成绩"
ws.append(["科目", "分数", "满分"])
ws.append(["语文", 96, 120])
ws.append(["数学", 82, 120])
ws2 = wb.create_sheet("计划")
ws2.append(["日期", "内容"])
ws2.append(["10-30", "复习光的折射"])
wb.save(os.path.join(out, "test.xlsx"))

print("生成完成:", out)
for name in ("test.pptx", "test.docx", "test.xlsx"):
    p = os.path.join(out, name)
    print(name, os.path.getsize(p), "bytes")
